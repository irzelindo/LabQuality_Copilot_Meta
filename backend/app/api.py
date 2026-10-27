import csv
import hashlib
import io
import secrets
import time
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Literal
from zoneinfo import ZoneInfo
from fastapi import APIRouter, Depends, File, Header, HTTPException, Query, Request, Response, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session
from .config import settings
from .database import get_db
from .models import *
from .schemas import *
from .security import *
from .services import *

def documented_client_header(x_labquality_client: str = Header(default="web", description="Envie web nos pedidos POST, PATCH, PUT e DELETE.")):
    return x_labquality_client


router = APIRouter(prefix="/api/v1", dependencies=[Depends(documented_client_header)])
DB = Depends(get_db)
ACTOR = Depends(current_user)
login_failures: dict[str, list[float]] = {}


def record(row, exclude=()) -> dict:
    result = {}
    for column in row.__table__.columns:
        if column.name in exclude:
            continue
        value = getattr(row, column.name)
        result[column.name] = value.isoformat() + "Z" if isinstance(value, datetime) else value
    return result


def today() -> str:
    return datetime.now(ZoneInfo("Africa/Maputo")).date().isoformat()


@router.post("/auth/login")
def login(body: Login, request: Request, response: Response, db: Session = DB):
    key = f"{request.client.host if request.client else 'unknown'}:{body.email.lower()}"
    recent = [t for t in login_failures.get(key, []) if t > time.monotonic() - 300]
    if len(recent) >= 8:
        raise HTTPException(429, "Demasiadas tentativas. Volte a tentar dentro de cinco minutos.")
    user = db.scalar(select(User).where(User.email == body.email.lower(), User.active.is_(True)))
    # Always perform an Argon2 verification, including unknown accounts.
    digest = user.password_hash if user else DUMMY_PASSWORD_HASH
    valid = passwords.verify(body.password, digest)
    if not user or not valid:
        login_failures[key] = recent + [time.monotonic()]
        raise HTTPException(401, "Email ou palavra-passe incorrectos.")
    login_failures.pop(key, None)
    token = create_session(db, user)
    response.set_cookie("lq_session", token, httponly=True, secure=settings().cookie_secure, samesite="strict",
                        max_age=settings().session_hours * 3600, path="/")
    audit(db, user, "auth.login", user)
    return user_out(user)


@router.post("/auth/logout")
def logout(request: Request, response: Response, db: Session = DB):
    token = request.cookies.get("lq_session")
    if token:
        db.execute(delete(LoginSession).where(LoginSession.token_hash == token_hash(token)))
    response.delete_cookie("lq_session", path="/")
    return {"ok": True}


@router.get("/auth/me")
def me(user: User = ACTOR):
    return user_out(user)


@router.post("/auth/password")
def change_password(body: PasswordChange, response: Response, db: Session = DB, user: User = ACTOR):
    if not passwords.verify(body.current_password, user.password_hash):
        raise HTTPException(400, "Palavra-passe actual incorrecta.")
    user.password_hash = passwords.hash(body.new_password)
    db.execute(delete(LoginSession).where(LoginSession.user_id == user.id))
    response.delete_cookie("lq_session", path="/")
    audit(db, user, "auth.password_changed", user)
    return {"ok": True}


@router.get("/laboratories")
def labs(db: Session = DB, user: User = ACTOR):
    return [record(l) for l in db.scalars(select(Laboratory).where(Laboratory.id.in_(laboratory_ids(db, user))).order_by(Laboratory.name))]


@router.post("/laboratories", status_code=201)
def add_lab(body: LabIn, db: Session = DB, user: User = ACTOR):
    role(user, {"admin"})
    lab = Laboratory(organization_id=user.organization_id, **body.model_dump())
    db.add(lab)
    db.flush()
    audit(db, user, "laboratory.created", lab)
    return record(lab)


@router.get("/users")
def users(db: Session = DB, user: User = ACTOR):
    require_details(user)
    allowed = set(laboratory_ids(db, user))
    rows = db.scalars(select(User).where(User.organization_id == user.organization_id).order_by(User.name))
    return [user_out(u) for u in rows if user.role == "admin" or u.id == user.id or (u.active and allowed.intersection(laboratory_ids(db, u)))]


@router.post("/users", status_code=201)
def add_user(body: UserIn, db: Session = DB, user: User = ACTOR):
    role(user, {"admin"})
    for lab_id in body.laboratory_ids:
        require_lab(db, user, lab_id)
    values = body.model_dump(exclude={"password"})
    row = User(organization_id=user.organization_id, password_hash=passwords.hash(body.password), **values)
    db.add(row)
    db.flush()
    audit(db, user, "user.created", row, {"role": row.role})
    return user_out(row)


@router.patch("/users/{user_id}/access")
def update_access(user_id: str, body: AccessUpdate, db: Session = DB, user: User = ACTOR):
    role(user, {"admin"})
    row = get_scoped(db, User, user_id, user)
    if row.id == user.id and (not body.active or body.role != "admin"):
        raise HTTPException(409, "Não pode remover o seu próprio acesso de administrador.")
    for lab_id in body.laboratory_ids:
        require_lab(db, user, lab_id)
    for k, v in body.model_dump().items():
        setattr(row, k, v)
    if not body.active:
        db.execute(delete(LoginSession).where(LoginSession.user_id == row.id))
    audit(db, user, "user.access_changed", row, body.model_dump())
    return user_out(row)


@router.get("/instruments")
def instruments(db: Session = DB, user: User = ACTOR):
    require_details(user)
    return [record(i) for i in db.scalars(select(Instrument).where(Instrument.organization_id == user.organization_id).order_by(Instrument.code, Instrument.version))]


@router.post("/instruments", status_code=201)
def add_instrument(body: InstrumentIn, db: Session = DB, user: User = ACTOR):
    role(user, {"admin", "manager"})
    row = Instrument(organization_id=user.organization_id, **body.model_dump())
    db.add(row)
    db.flush()
    audit(db, user, "instrument.published", row, {"scoring_rule": row.scoring_rule})
    return record(row)


@router.post("/instruments/xlsform")
async def import_pilot_xlsform(file: UploadFile = File(...), dry_run: bool = True,
                               db: Session = DB, user: User = ACTOR):
    """Recognize the audited SLIPTA XLSForm, without guessing scoring for other forms."""
    from .demo import PILOT
    role(user, {"admin", "manager"})
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(415, "Seleccione o XLSForm em formato XLSX.")
    content = await file.read(2_000_001)
    if len(content) > 2_000_000:
        raise HTTPException(413, "Limite do XLSForm: 2 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30_000_000:
                raise HTTPException(413, "O XLSForm expandido excede o limite.")
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            if not {"survey", "choices", "settings"}.issubset(workbook.sheetnames):
                raise HTTPException(422, "Faltam as folhas survey, choices ou settings do XLSForm.")
            rows = workbook["settings"].iter_rows(min_row=1, max_row=2, values_only=True)
            headers, values = next(rows), next(rows)
            metadata = dict(zip(headers, values))
        finally:
            workbook.close()
    except (zipfile.BadZipFile, ValueError, KeyError, StopIteration, OSError) as error:
        raise HTTPException(422, "Não foi possível ler o XLSForm.") from error
    form_id = str(metadata.get("form_id") or "")
    if form_id != PILOT["form_id"] or hashlib.sha256(content).hexdigest() != PILOT["source_sha256"]:
        raise HTTPException(422, "Este XLSForm ou esta revisão ainda não tem mapeamento de pontuação validado. Não foi criado um instrumento.")
    criteria = [{key: item[key] for key in ("code", "section", "title", "maximum", "critical")}
                for item in PILOT["criteria"]]
    existing = db.scalar(select(Instrument).where(Instrument.organization_id == user.organization_id,
        Instrument.code == PILOT["instrument_code"], Instrument.version == PILOT["instrument_version"]))
    if existing and (existing.criteria != criteria or existing.scoring_rule != "sum_applicable_v1"):
        raise HTTPException(409, "A versão do instrumento já existe com critérios diferentes. Reveja-a antes de importar.")
    summary = {"form_id": form_id, "form_version": str(metadata.get("version") or ""),
        "instrument_code": PILOT["instrument_code"], "instrument_version": PILOT["instrument_version"],
        "criteria_count": len(criteria), "maximum": sum(c["maximum"] for c in criteria),
        "existing": bool(existing), "dry_run": dry_run}
    if dry_run or existing:
        return {**summary, "instrument_id": existing.id if existing else None,
                "message": "O instrumento já existe nesta instituição." if existing else "XLSForm reconhecido; pode criar o instrumento piloto."}
    row = Instrument(organization_id=user.organization_id, code=PILOT["instrument_code"],
        version=PILOT["instrument_version"], name="Checklist SLIPTA ODK · piloto",
        description="151 itens/365 pontos do XLSForm validado. Pontuação provisória; sem classificação oficial SLIPTA.",
        scoring_rule="sum_applicable_v1", criteria=criteria)
    db.add(row)
    db.flush()
    audit(db, user, "instrument.published", row, {"source": "xlsform", "form_id": form_id})
    return {**summary, "instrument_id": row.id, "message": "Instrumento piloto criado a partir do XLSForm validado."}


def assessment_out(db: Session, row: Assessment, detail=False):
    instrument = db.get(Instrument, row.instrument_id)
    root = db.get(AssessmentRoot, row.root_id)
    result = record(row)
    result.update(laboratory_name=db.get(Laboratory, row.laboratory_id).name,
                  laboratory_code=db.get(Laboratory, row.laboratory_id).code,
                  instrument_name=instrument.name, instrument_code=instrument.code, instrument_version=instrument.version,
                  source=root.source, external_id=root.external_id, current=row.revision == root.head_revision, score=score(row, instrument))
    if detail:
        result["criteria"] = instrument.criteria
        result["findings"] = [record(f) for f in db.scalars(select(Finding).where(Finding.assessment_id == row.id))]
        result["revisions"] = [{"id": a.id, "revision": a.revision, "reason": a.reason, "status": a.status}
                               for a in db.scalars(select(Assessment).where(Assessment.root_id == row.root_id).order_by(Assessment.revision.desc()))]
    else:
        result.pop("answers", None)
    return result


@router.get("/assessments")
def assessments(laboratory_id: str | None = None, start: date | None = None, end: date | None = None,
                origin: Literal["all", "odk", "demo"] = "all", offset: int = Query(0, ge=0),
                limit: int = Query(100, ge=1, le=500), db: Session = DB, user: User = ACTOR):
    require_details(user)
    rows = filter_assessments_by_origin(db,
        current_assessments(db, user, laboratory_id, start.isoformat() if start else None, end.isoformat() if end else None), origin)
    return {"items": [assessment_out(db, a) for a in rows[offset:offset + limit]], "total": len(rows)}


@router.get("/assessments/{assessment_id}")
def assessment_detail(assessment_id: str, db: Session = DB, user: User = ACTOR):
    require_details(user)
    return assessment_out(db, get_scoped(db, Assessment, assessment_id, user), True)


@router.post("/imports/assessments")
def import_json(body: AssessmentBatch, dry_run: bool = False, db: Session = DB, user: User = ACTOR):
    role(user, EDITORS)
    if dry_run:
        transaction = db.begin_nested()
        result = import_assessments(db, user, body.assessments)
        transaction.rollback()
        return {**result, "assessment_ids": [], "import_run_id": None, "dry_run": True}
    return import_assessments(db, user, body.assessments)


def parse_table(rows: list[dict]) -> list[AssessmentIn]:
    groups = {}
    needed = {"source", "external_id", "laboratory_code", "instrument_code", "instrument_version", "assessment_date", "criterion_code", "value"}
    if not rows or not needed.issubset(rows[0]):
        raise HTTPException(422, "Colunas obrigatórias: " + ", ".join(sorted(needed)))
    for i, row in enumerate(rows):
        key = (str(row["source"]), str(row["external_id"]))
        meta = {k: str(row[k])[:200] for k in ("source", "external_id", "laboratory_code", "instrument_code", "instrument_version")}
        d = row["assessment_date"]
        meta["assessment_date"] = d.date().isoformat() if isinstance(d, datetime) else str(d)
        if key not in groups:
            groups[key] = {**meta, "source_revision": str(row.get("source_revision") or ""), "answers": {}}
        if any(groups[key][k] != v for k, v in meta.items()):
            raise HTTPException(422, f"Metadados incompatíveis para a mesma avaliação, linha {i + 2}.")
        code = str(row["criterion_code"])
        if code in groups[key]["answers"]:
            raise HTTPException(422, f"Critério repetido, linha {i + 2}.")
        raw = row["value"]
        is_na = isinstance(raw, str) and raw.strip().upper() in {"NA", "N/A"}
        value = None if is_na or raw is None or str(raw).strip() == "" else raw
        groups[key]["answers"][code] = {"value": value, "comment": str(row.get("comment") or ""),
                                         "not_applicable": is_na}
    try:
        return [AssessmentIn.model_validate(v) for v in groups.values()]
    except ValueError as error:
        raise HTTPException(422, "Dados inválidos no ficheiro. Confirme datas ISO e pontuações inteiras.") from error


@router.post("/imports/file")
async def import_file(file: UploadFile = File(...), dry_run: bool = False, db: Session = DB, user: User = ACTOR):
    role(user, EDITORS)
    content = await file.read(2_000_001)
    if len(content) > 2_000_000:
        raise HTTPException(413, "Limite de importação: 2 MB.")
    filename = (file.filename or "").lower()
    try:
        if filename.endswith(".csv"):
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > 10000:
                    raise HTTPException(413, "Limite: 10 000 linhas por ficheiro.")
        elif filename.endswith(".xlsx"):
            import openpyxl
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 30_000_000:
                    raise HTTPException(413, "O Excel expandido excede o limite.")
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                if {"survey", "choices", "settings"}.issubset(workbook.sheetnames):
                    raise HTTPException(422, "Este ficheiro é um XLSForm, não uma tabela de avaliações. Abra Instrumentos → Carregar XLSForm; para ler respostas use a ligação ODK Central.")
                sheet = workbook.active
                iterator = sheet.iter_rows(values_only=True)
                headers = [str(v) for v in next(iterator)]
                rows = []
                for values in iterator:
                    if any(v is not None for v in values):
                        rows.append(dict(zip(headers, values)))
                    if len(rows) > 10000:
                        raise HTTPException(413, "Limite: 10 000 linhas por ficheiro.")
            finally:
                workbook.close()
        else:
            raise HTTPException(415, "Utilize CSV UTF-8 ou Excel XLSX.")
        records = parse_table(rows)
    except (UnicodeError, zipfile.BadZipFile, StopIteration, KeyError, TypeError) as error:
        raise HTTPException(422, "Não foi possível ler o ficheiro.") from error
    if len(records) > 500:
        raise HTTPException(413, "Limite: 500 avaliações por importação.")
    return import_json(AssessmentBatch(assessments=records), dry_run, db, user)


@router.post("/assessments/{assessment_id}/finalize")
def finalize(assessment_id: str, body: VersionIn, db: Session = DB, user: User = ACTOR):
    role(user, APPROVERS)
    row = get_scoped(db, Assessment, assessment_id, user)
    ensure_version(row, body.version)
    if row.revision != db.get(AssessmentRoot, row.root_id).head_revision or row.status != "review":
        raise HTTPException(409, "Só a revisão actual em análise pode ser finalizada.")
    if not score(row, db.get(Instrument, row.instrument_id))["complete"]:
        raise HTTPException(422, "Complete as respostas antes de finalizar.")
    row.status = "finalized"
    audit(db, user, "assessment.finalized", row)
    db.flush()
    return assessment_out(db, row, True)


@router.get("/findings")
def findings(laboratory_id: str | None = None, db: Session = DB, user: User = ACTOR):
    require_details(user)
    assessment_rows = current_assessments(db, user, laboratory_id)
    query = select(Finding).where(Finding.assessment_id.in_([a.id for a in assessment_rows])).order_by(Finding.critical.desc(), Finding.created_at.desc()).limit(1000)
    return [record(f) for f in db.scalars(query)]


def action_out(db: Session, action: Action, detail=False):
    result = record(action)
    result.update(laboratory_name=db.get(Laboratory, action.laboratory_id).name,
                  owner_name=db.get(User, action.owner_id).name,
                  overdue=action.status != "closed" and action.due_date < today())
    if detail:
        result["evidence"] = [record(e, {"storage_key"}) for e in db.scalars(select(Evidence).where(Evidence.action_id == action.id))]
        result["verifications"] = [record(v) for v in db.scalars(select(Verification).where(Verification.action_id == action.id).order_by(Verification.created_at.desc()))]
        result["finding"] = record(db.get(Finding, action.finding_id))
    return result


def create_action(body: ActionIn, db: Session, user: User):
    finding = get_scoped(db, Finding, body.finding_id, user)
    eligible_owner(db, user, body.owner_id, finding.laboratory_id)
    values = body.model_dump()
    values["due_date"] = body.due_date.isoformat()
    row = Action(organization_id=user.organization_id, laboratory_id=finding.laboratory_id, creator_id=user.id, **values)
    db.add(row)
    db.flush()
    audit(db, user, "action.created", row)
    notify(db, user, row.owner_id, row.laboratory_id, f"action:{row.id}:assigned:1", row.title, "Foi-lhe atribuída uma proposta de acção.")
    return row


@router.post("/actions", status_code=201)
def add_action(body: ActionIn, db: Session = DB, user: User = ACTOR):
    role(user, EDITORS)
    return action_out(db, create_action(body, db, user), True)


@router.get("/actions")
def actions(laboratory_id: str | None = None, status: str | None = None, db: Session = DB, user: User = ACTOR):
    require_details(user)
    labs = laboratory_ids(db, user)
    if laboratory_id:
        require_lab(db, user, laboratory_id)
        labs = [laboratory_id]
    query = select(Action).where(Action.organization_id == user.organization_id, Action.laboratory_id.in_(labs))
    if status:
        query = query.where(Action.status == status)
    return [action_out(db, a) for a in db.scalars(query.order_by(Action.due_date).limit(1000))]


@router.get("/actions/{action_id}")
def action_detail(action_id: str, db: Session = DB, user: User = ACTOR):
    require_details(user)
    return action_out(db, get_scoped(db, Action, action_id, user), True)


@router.patch("/actions/{action_id}")
def edit_action(action_id: str, body: ActionEdit, db: Session = DB, user: User = ACTOR):
    role(user, EDITORS)
    row = get_scoped(db, Action, action_id, user)
    ensure_version(row, body.version)
    if row.status not in {"proposed", "approved", "in_progress", "reopened"}:
        raise HTTPException(409, "A acção já está em verificação ou encerrada.")
    eligible_owner(db, user, body.owner_id, row.laboratory_id)
    previous_owner = row.owner_id
    for key, value in body.model_dump(exclude={"version"}).items():
        setattr(row, key, value.isoformat() if isinstance(value, date) else value)
    audit(db, user, "action.updated", row)
    db.flush()
    if previous_owner != row.owner_id:
        notify(db, user, row.owner_id, row.laboratory_id, f"action:{row.id}:assigned:{row.version}", row.title, "A acção foi-lhe atribuída.")
    return action_out(db, row, True)


@router.post("/actions/{action_id}/transition")
def action_transition(action_id: str, body: StateIn, db: Session = DB, user: User = ACTOR):
    row = get_scoped(db, Action, action_id, user)
    ensure_version(row, body.version)
    allowed = {"proposed": {"approved"}, "approved": {"in_progress"}, "in_progress": {"in_verification"},
               "in_verification": {"reopened"}, "closed": {"reopened"}, "reopened": {"in_progress"}}
    if body.status not in allowed.get(row.status, set()):
        raise HTTPException(409, "Transição de estado inválida.")
    role(user, APPROVERS if body.status in {"approved", "reopened"} else EDITORS)
    if body.status == "reopened" and len(body.comment) < 5:
        raise HTTPException(422, "Indique o motivo da reabertura.")
    if body.status == "in_verification" and not db.scalar(select(Evidence.id).where(Evidence.action_id == row.id).limit(1)):
        raise HTTPException(422, "Anexe uma evidência antes de solicitar verificação.")
    previous = row.status
    row.status = body.status
    row.closed_at = None
    audit(db, user, "action.transition", row, {"from": previous, "to": body.status, "comment": body.comment})
    db.flush()
    return action_out(db, row, True)


@router.post("/actions/{action_id}/verify")
def verify_action(action_id: str, body: VerifyIn, db: Session = DB, user: User = ACTOR):
    role(user, APPROVERS)
    row = get_scoped(db, Action, action_id, user)
    ensure_version(row, body.version)
    if row.status != "in_verification":
        raise HTTPException(409, "A acção não está em verificação.")
    if row.owner_id == user.id:
        raise HTTPException(403, "A verificação deve ser feita por outra pessoa.")
    for evidence_id in body.evidence_ids:
        evidence = get_scoped(db, Evidence, evidence_id, user)
        if evidence.action_id != row.id:
            raise HTTPException(422, "A evidência não pertence a esta acção.")
    verification = Verification(organization_id=user.organization_id, action_id=row.id, verifier_id=user.id,
                                result=body.result, comment=body.comment, evidence_ids=body.evidence_ids)
    db.add(verification)
    row.status = "closed" if body.result == "effective" else "reopened"
    row.closed_at = now() if row.status == "closed" else None
    audit(db, user, "action.verified", row, {"result": body.result, "comment": body.comment, "evidence_ids": body.evidence_ids})
    notify(db, user, row.owner_id, row.laboratory_id, f"action:{row.id}:verification:{row.version}", row.title, "Resultado da verificação: " + body.result)
    db.flush()
    return action_out(db, row, True)


@router.post("/actions/{action_id}/evidence", status_code=201)
async def add_evidence(action_id: str, file: UploadFile = File(...), note: str = "", db: Session = DB, user: User = ACTOR):
    role(user, EDITORS | {"verifier"})
    action = get_scoped(db, Action, action_id, user)
    if action.status == "closed":
        raise HTTPException(409, "Reabra a acção antes de acrescentar evidências.")
    content = await file.read(10_000_001)
    if not content or len(content) > 10_000_000:
        raise HTTPException(413, "O ficheiro deve ter entre 1 byte e 10 MB.")
    original = (file.filename or "evidence").replace("\\", "/").split("/")[-1]
    extension = Path(original).suffix.lower()
    signatures = {".pdf": (b"%PDF-", "application/pdf"), ".png": (b"\x89PNG\r\n\x1a\n", "image/png"),
                  ".jpg": (b"\xff\xd8\xff", "image/jpeg"), ".jpeg": (b"\xff\xd8\xff", "image/jpeg")}
    if extension == ".txt":
        try:
            content.decode("utf-8")
        except UnicodeError as error:
            raise HTTPException(415, "O texto deve estar em UTF-8.") from error
        media_type = "text/plain"
    elif extension in signatures and content.startswith(signatures[extension][0]):
        media_type = signatures[extension][1]
    else:
        raise HTTPException(415, "Utilize PDF, PNG, JPEG ou TXT com conteúdo válido.")
    key = secrets.token_hex(24)
    directory = Path(settings().storage_dir).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / key
    with path.open("xb") as handle:
        handle.write(content)
    path.chmod(0o600)
    db.info.setdefault("created_files", []).append(path)
    row = Evidence(organization_id=user.organization_id, laboratory_id=action.laboratory_id, action_id=action.id,
                   filename=original[:255], media_type=media_type, storage_key=key, size_bytes=len(content),
                   sha256=hashlib.sha256(content).hexdigest(), note=note[:5000], author_id=user.id)
    db.add(row)
    db.flush()
    audit(db, user, "evidence.added", row, {"sha256": row.sha256})
    return record(row, {"storage_key"})


@router.get("/evidence/{evidence_id}/download")
def download_evidence(evidence_id: str, db: Session = DB, user: User = ACTOR):
    require_details(user)
    row = get_scoped(db, Evidence, evidence_id, user)
    path = Path(settings().storage_dir).resolve() / row.storage_key
    if not path.is_file():
        raise HTTPException(404, "Ficheiro indisponível.")
    audit(db, user, "evidence.downloaded", row)
    return FileResponse(path, media_type=row.media_type, filename=row.filename, headers={"X-Content-Type-Options": "nosniff"})
