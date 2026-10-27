import hashlib
import io

from openpyxl import Workbook
from sqlalchemy import select

from app.database import session_factory
from app.demo import PILOT
from app.models import Instrument
from conftest import login


def xlsform_bytes():
    workbook = Workbook()
    survey = workbook.active
    survey.title = "survey"
    survey.append(["type", "name", "label"])
    workbook.create_sheet("choices").append(["list_name", "name", "label"])
    settings = workbook.create_sheet("settings")
    settings.append(["form_title", "form_id", "version"])
    settings.append(["SLIPTA", PILOT["form_id"], "2026090801"])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def post_form(client, path, content):
    return client.post(path, files={"file": ("slipta.xlsx", content,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_xlsform_is_not_an_assessment_file(client):
    login(client)
    response = post_form(client, "/api/v1/imports/file?dry_run=true", xlsform_bytes())
    assert response.status_code == 422
    assert "Instrumentos" in response.json()["detail"]


def test_verified_xlsform_creates_one_instrument_without_assessments(client, monkeypatch):
    content = xlsform_bytes()
    monkeypatch.setitem(PILOT, "source_sha256", hashlib.sha256(content).hexdigest())
    login(client)
    path = "/api/v1/instruments/xlsform"
    preview = post_form(client, path + "?dry_run=true", content)
    assert preview.status_code == 200, preview.text
    assert preview.json()["criteria_count"] == 151
    assert preview.json()["maximum"] == 365
    assert preview.json()["form_id"] == PILOT["form_id"]
    assert preview.json()["form_version"] == "2026090801"
    assert not preview.json()["existing"]
    with session_factory()() as db:
        assert not db.scalar(select(Instrument).where(Instrument.code == PILOT["instrument_code"]))
    created = post_form(client, path + "?dry_run=false", content)
    assert created.status_code == 200, created.text
    assert created.json()["instrument_id"]
    again = post_form(client, path + "?dry_run=false", content)
    assert again.status_code == 200
    assert again.json()["existing"]
    with session_factory()() as db:
        rows = list(db.scalars(select(Instrument).where(Instrument.code == PILOT["instrument_code"])))
        assert len(rows) == 1 and rows[0].published
    login(client, "qualidade@labquality.local")
    assert post_form(client, path, content).status_code == 403


def test_unverified_xlsform_cannot_create_instrument(client):
    login(client)
    response = post_form(client, "/api/v1/instruments/xlsform?dry_run=false", xlsform_bytes())
    assert response.status_code == 422
    assert "mapeamento de pontuação" in response.json()["detail"]
