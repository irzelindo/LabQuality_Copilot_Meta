import hashlib
import io

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from app.backup import create_backup
from app.config import settings
from app.database import session_factory
from app.demo import PILOT
from app.import_odk_pilot import run
from app.models import AssessmentRoot, Instrument, Laboratory, User
from app.schemas import AssessmentIn


def test_local_odk_import_previews_backs_up_and_does_not_duplicate(client, monkeypatch, tmp_path, capsys):
    from app import import_odk_pilot as importer

    cfg = settings()
    cfg.odk_url, cfg.odk_email, cfg.odk_password = "https://odk.example", "test@example.invalid", "test-only"
    workbook = Workbook()
    workbook.active.title = "survey"
    workbook.create_sheet("choices")
    metadata = workbook.create_sheet("settings")
    metadata.append(["form_title", "form_id", "version"])
    metadata.append(["SLIPTA", PILOT["form_id"], "2026090801"])
    content = io.BytesIO()
    workbook.save(content)
    form_file = tmp_path / "slipta.xlsx"
    form_file.write_bytes(content.getvalue())
    monkeypatch.setitem(PILOT, "source_sha256", hashlib.sha256(content.getvalue()).hexdigest())

    records = [AssessmentIn(source=f"odk:1:{PILOT['form_id']}", external_id=f"uuid:test-{i}",
        laboratory_code="LAB01" if i < 5 else "LAB02", instrument_code=PILOT["instrument_code"],
        instrument_version=PILOT["instrument_version"], assessment_date="2026-09-19",
        answers={"1.1": {"value": 1, "comment": "Observação do teste."}}) for i in range(9)]
    laboratory_metadata = [{"laboratory_name": "Laboratório Um" if i < 5 else "Laboratório Dois"} for i in range(9)]
    monkeypatch.setattr(importer, "fetch_odk_with_metadata", lambda _: (records, laboratory_metadata))
    backups = []

    def backup(_):
        created = create_backup(tmp_path / "backups")
        backups.append(created)
        return created

    monkeypatch.setattr(importer, "create_backup", backup)
    with pytest.raises(ValueError, match="Instrumento SLIPTA-ODK-PILOT 2026.10 ausente"):
        run(1, 9, "admin@labquality.local", apply=False, register_missing=False)
    run(1, 9, "admin@labquality.local", apply=False, register_missing=False, xlsform_path=form_file)
    assert "por cadastrar" in capsys.readouterr().out
    assert backups == []
    with session_factory()() as db:
        assert db.scalar(select(func.count()).select_from(AssessmentRoot).where(AssessmentRoot.source == records[0].source)) == 0
        assert db.scalar(select(func.count()).select_from(Laboratory).where(Laboratory.code.like("LAB0%"))) == 0
        assert db.scalar(select(func.count()).select_from(Instrument).where(Instrument.code == PILOT["instrument_code"])) == 0

    run(1, 9, "admin@labquality.local", apply=True, register_missing=True, xlsform_path=form_file)
    assert "9 novas" in capsys.readouterr().out
    assert len(backups) == 1 and (backups[0] / "database.sqlite").is_file()
    run(1, 9, "admin@labquality.local", apply=True, register_missing=False)
    assert "9 inalteradas" in capsys.readouterr().out
    with session_factory()() as db:
        assert db.scalar(select(func.count()).select_from(AssessmentRoot).where(AssessmentRoot.source == records[0].source)) == 9
        assert db.scalar(select(func.count()).select_from(Laboratory).where(Laboratory.code.like("LAB0%"))) == 2
        assert db.scalar(select(func.count()).select_from(Instrument).where(Instrument.code == PILOT["instrument_code"])) == 1


def test_local_odk_import_never_overwrites_incompatible_instrument(client):
    from app.import_odk_pilot import check_scope

    with session_factory()() as db:
        admin = db.scalar(select(User).where(User.email == "admin@labquality.local"))
        instrument = Instrument(organization_id=admin.organization_id, code=PILOT["instrument_code"],
            version=PILOT["instrument_version"], name="Outra versão", scoring_rule="sum_applicable_v1",
            criteria=[{"code": "1.1", "section": "Teste", "title": "Divergente", "maximum": 5, "critical": False}])
        db.add(instrument)
        db.commit()
        with pytest.raises(ValueError, match="1 critérios existentes; esperados 151"):
            check_scope(db, admin, {})
        assert db.get(Instrument, instrument.id).criteria[0]["maximum"] == 5


def test_local_odk_import_refuses_changed_count(client, monkeypatch):
    from app import import_odk_pilot as importer

    cfg = settings()
    cfg.odk_url, cfg.odk_email, cfg.odk_password = "https://odk.example", "test@example.invalid", "test-only"
    monkeypatch.setattr(importer, "fetch_odk_with_metadata", lambda _: ([], []))
    with pytest.raises(ValueError, match="Esperadas 9 submissões"):
        run(1, 9, "admin@labquality.local", apply=True, register_missing=True)
